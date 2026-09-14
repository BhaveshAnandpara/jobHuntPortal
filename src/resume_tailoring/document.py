"""Renders a `ResumeProfile` as a clean, ATS-friendly `.docx` document.

Nothing else in this codebase writes `.docx` files — `src/profiles/parsing.py`
only reads them — so this module is the single place that knows the tailored
resume's layout. It is deliberately dumb: it takes an already-tailored
`ResumeProfile` (Task 3's `TailoringResult.tailored_profile`) and renders it.
It makes no LLM calls, reads no database, and touches no filesystem.

**Interface contract (for the storage layer, TECH.md Task 5).**
`build_resume_document(profile)` returns the finished `.docx` file as `bytes`,
ready to hand straight to `LocalTailoredResumeStorage.save(..., content=...)`
(the same `content: bytes` shape `src/profiles/storage.py::LocalResumeStorage`
already takes) or to a streaming download response. Nothing is written to
disk here, so the caller owns naming, placement, and lifetime of the file.

**Visual reference.** The palette, typography, and paragraph geometry below
are lifted from the user-supplied reference resume (`template.docx`), read
back run-by-run rather than eyeballed: Calibri throughout; three colors
(`1A1A1A` ink, `1F4E5F` accent, `555555` muted); a centered header whose
name is 16pt bold, whose contact line is 9pt muted, and whose links line
carries real external hyperlinks; uppercase accent-colored section headings
with a hairline bottom border; a right-aligned tab stop carrying dates on
role and education lines; and a bold lead-in label on the skills line.

**ATS constraints (PRODUCT.md "Document Content & Format").** The document is
a flat sequence of headings and paragraphs only: no tables used for layout,
no text boxes, no headers/footers, no images — all of which ATS parsers
commonly mishandle. Section titles still use python-docx's built-in
`Heading 1`/`Heading 2` styles (restyled, not replaced by hand-bolded
paragraphs) so the file keeps a real semantic outline, and bullet points
still use the built-in `List Bullet` style rather than typed-in dashes. The
template's uppercase headings are reproduced with `w:caps`, which uppercases
only the *rendering* — the underlying text stays "Work Experience", so a
parser reading the XML sees the ordinary section name.

**Profession independence (CLAUDE.md rules 3-4).** Every section is driven
purely by the profile's own content — no industry-specific vocabulary,
ordering, or assumptions live here. A section whose source data is empty is
omitted entirely, heading included, rather than rendered as an empty stub.
The same holds for the three profile links: each is rendered only when the
profile actually carries it, and a missing one leaves no label, no
placeholder, and no dangling separator behind.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Length, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from shared.types.dto import (
    ContactInfo,
    EducationEntry,
    ResumeProfile,
    WorkExperienceEntry,
)

# Rendering vocabulary, named so tests (and any future caller) assert against
# the constant rather than a duplicated literal.
NAME_FALLBACK = "Resume"
"""Used as the document's heading when `contact.full_name` is unavailable.

Should be rare after Task 2's enrichment, but a missing name must never
crash generation or produce a document headed "None".
"""

CURRENT_ROLE_LABEL = "Present"
"""Shown in place of an end date when `WorkExperienceEntry.end_date is None`
(the DTO's documented encoding of a current role)."""

CONTACT_SEPARATOR = "   |   "
DATE_SEPARATOR = " – "  # en dash
TITLE_COMPANY_SEPARATOR = " — "  # em dash
SKILLS_LABEL = "Skills: "
"""The bold lead-in on the skills line.

`ResumeProfile.skills` is a flat `list[str]`, so this is one label over one
comma-separated line. The reference template splits skills into several
labelled groups ("AI Agent Systems: ", "Full-Stack: ", ...), which would
require `CandidateProfile.skills` to become grouped data — a schema change
well outside this module, noted rather than invented here.
"""

SECTION_SUMMARY = "Summary"
SECTION_SKILLS = "Skills"
SECTION_EXPERIENCE = "Work Experience"
SECTION_PROJECTS = "Projects"
SECTION_EDUCATION = "Education"
SECTION_CERTIFICATIONS = "Certifications"

LINKEDIN_LABEL = "LinkedIn"
GITHUB_LABEL = "GitHub"
PORTFOLIO_LABEL = "Portfolio"

_NAME_HEADING_LEVEL = 1
_SECTION_HEADING_LEVEL = 2
_BULLET_STYLE = "List Bullet"

# -- palette (template.docx) -------------------------------------------------
INK_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT_COLOR = RGBColor(0x1F, 0x4E, 0x5F)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)
ACCENT_HEX = "1F4E5F"
"""The accent as OOXML writes it — needed for the raw `w:pBdr` element,
which takes a hex string rather than an `RGBColor`."""

FONT_NAME = "Calibri"

# -- type scale (template.docx, in half-points there, whole points here) -----
NAME_SIZE = Pt(16)
CONTACT_SIZE = Pt(9)
SECTION_HEADING_SIZE = Pt(10)
BODY_SIZE = Pt(9.5)

_BORDER_WIDTH_EIGHTHS = "6"  # w:sz is in eighths of a point -> 0.75pt hairline
_BORDER_SPACE_POINTS = "2"

# CT_PPr's element order (OOXML CT_PPr sequence). `w:pBdr` is not modelled by
# python-docx, so it is inserted by hand *before* everything that may legally
# follow it — appending it blindly would produce a schema-invalid ordering
# that Word repairs (and, while repairing, may discard).
_AFTER_P_BORDERS = (
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)


# ---------------------------------------------------------------------------
# Small text helpers
# ---------------------------------------------------------------------------


def _clean(value: str | None) -> str:
    """Normalize an optional free-text field to a trimmed string ("" when
    absent), so "present but blank" and "absent" render identically.
    """
    return value.strip() if value else ""


def _clean_all(values: Iterable[str] | None) -> list[str]:
    """Trim a list of free-text values and drop the empties."""
    return [text for text in (_clean(value) for value in values or []) if text]


def _join(parts: Iterable[str | None], separator: str) -> str:
    """Join only the parts that actually have content."""
    return separator.join(part for part in (_clean(p) for p in parts) if part)


def _date_range(entry: WorkExperienceEntry) -> str:
    """Render one entry's dates, never leaking a literal "None".

    A missing `end_date` means the role is current, so it renders as
    `CURRENT_ROLE_LABEL`. A missing `start_date` (resumes sometimes omit it)
    degrades to just the end side rather than a dangling separator.
    """
    start = _clean(entry.start_date)
    end = _clean(entry.end_date) or CURRENT_ROLE_LABEL
    return f"{start}{DATE_SEPARATOR}{end}" if start else end


def _education_line(entry: EducationEntry) -> str:
    """"BSc, Adult Nursing" — whichever of degree/field of study exist."""
    return _join([entry.degree, entry.field_of_study], ", ")


def _institution_line(entry: EducationEntry) -> str:
    """The awarding body on its own; the year is tab-aligned separately."""
    return _clean(entry.institution)


def _graduation_year(entry: EducationEntry) -> str:
    return str(entry.graduation_year) if entry.graduation_year is not None else ""


# ---------------------------------------------------------------------------
# Selective emphasis
# ---------------------------------------------------------------------------


def _skill_pattern(skills: Sequence[str]) -> re.Pattern[str] | None:
    """A case-insensitive, whole-word alternation over the profile's skills.

    Longest-first so "Critical care nursing" wins over a bare "nursing"
    that also happens to be listed. Word boundaries are only applied where
    they mean anything: a term that starts or ends with punctuation (e.g.
    "C++", ".NET") gets none on that side, because `\\b` next to a
    non-word character matches in exactly the wrong places.
    """
    terms = sorted(set(_clean_all(skills)), key=len, reverse=True)
    if not terms:
        return None
    alternatives = []
    for term in terms:
        left = r"\b" if term[0].isalnum() or term[0] == "_" else ""
        right = r"\b" if term[-1].isalnum() or term[-1] == "_" else ""
        alternatives.append(f"{left}{re.escape(term)}{right}")
    return re.compile("|".join(alternatives), re.IGNORECASE)


def emphasis_segments(
    text: str, pattern: re.Pattern[str] | None
) -> list[tuple[str, bool]]:
    """Split `text` into `(fragment, emphasized)` pairs.

    Purely presentational: the fragments concatenate back to `text`
    character-for-character, so this pass can only ever change which runs
    carry `w:b` — never a single character of what the resume says. That
    property is what makes bolding safe to do after the tailoring LLM has
    finished, rather than asking the model to mark up its own output.
    """
    if pattern is None or not text:
        return [(text, False)]
    segments: list[tuple[str, bool]] = []
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            segments.append((text[cursor : match.start()], False))
        segments.append((match.group(0), True))
        cursor = match.end()
    if cursor < len(text):
        segments.append((text[cursor:], False))
    return segments or [(text, False)]


# ---------------------------------------------------------------------------
# OOXML primitives python-docx has no high-level API for
# ---------------------------------------------------------------------------


def _style_run(
    run: Run,
    *,
    size: Length,
    color: RGBColor,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
    caps: bool = False,
) -> Run:
    """Apply the template's run formatting explicitly.

    Every property is written explicitly — including the false ones — so a
    run never inherits bold/italic from a style and so the "this fragment is
    emphasized, that one is not" distinction is visible in the XML rather
    than implied by absence.
    """
    font = run.font
    font.name = FONT_NAME
    font.size = size
    font.color.rgb = color
    font.bold = bold
    font.italic = italic
    font.underline = underline

    rpr = run._element.get_or_add_rPr()
    # `font.name` only sets w:ascii/w:hAnsi; the reference file also pins the
    # complex-script and East Asian faces, without which Word substitutes a
    # different font for e.g. an accented or CJK character mid-line.
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:cs"), FONT_NAME)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)
    if bold:
        rpr.get_or_add_bCs().val = True
    if italic:
        rpr.get_or_add_iCs().val = True
    if caps:
        rpr.get_or_add_caps().val = True
    return run


def _add_text(
    paragraph: Paragraph, text: str, **formatting: object
) -> Run:
    return _style_run(paragraph.add_run(text), **formatting)  # type: ignore[arg-type]


def _add_bottom_border(paragraph: Paragraph) -> None:
    """The accent hairline under a section heading (`w:pBdr/w:bottom`)."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), _BORDER_WIDTH_EIGHTHS)
    bottom.set(qn("w:space"), _BORDER_SPACE_POINTS)
    bottom.set(qn("w:color"), ACCENT_HEX)
    borders.append(bottom)
    ppr.insert_element_before(borders, *_AFTER_P_BORDERS)


def add_external_hyperlink(paragraph: Paragraph, text: str, url: str) -> Run:
    """Append `text` to `paragraph` as a real, clickable external link.

    python-docx exposes no high-level hyperlink API, so this builds the
    `w:hyperlink` element itself and registers the target as an *external*
    package relationship (`part.relate_to(..., is_external=True)`) — the
    documented pattern, and the only one that survives a round trip through
    Word. The returned `Run` is the link's own run, so the caller styles it
    through the ordinary font API rather than more raw XML.

    Nothing here invents a URL: callers pass a URL the profile actually
    carries, or do not call this at all.
    """
    relationship_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")
    run_element = OxmlElement("w:r")
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)

    run = Run(run_element, paragraph)
    run.text = text
    return run


# ---------------------------------------------------------------------------
# Document primitives
# ---------------------------------------------------------------------------


def _tab_position(document: DocumentObject) -> Length:
    """The x-offset of the right margin — where a right tab stop belongs.

    Derived from the section rather than hard-coded to the reference file's
    9360 twips, which was measured against different margins; computing it
    keeps dates flush with the right edge whatever page size is in force.
    """
    section = document.sections[0]
    width = section.page_width or Inches(8.5)
    left = section.left_margin or Emu(0)
    right = section.right_margin or Emu(0)
    return Emu(int(width - left - right))


def _new_document() -> DocumentObject:
    """A blank document restyled to the reference resume's look.

    Only style *properties* are adjusted — the style names, and therefore
    the document's semantic outline, are left exactly as python-docx defines
    them, so `Heading 1`/`Heading 2`/`List Bullet` still mean what they mean
    to anything reading the file structurally.
    """
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    body = document.styles["Normal"]
    body.font.name = FONT_NAME
    body.font.size = BODY_SIZE
    body.font.color.rgb = INK_COLOR
    body.paragraph_format.space_after = Pt(2)

    name_heading = document.styles[f"Heading {_NAME_HEADING_LEVEL}"]
    name_heading.font.name = FONT_NAME
    name_heading.font.color.rgb = INK_COLOR
    name_heading.font.size = NAME_SIZE
    name_heading.font.bold = True
    name_heading.paragraph_format.space_before = Pt(0)
    name_heading.paragraph_format.space_after = Pt(2)

    section_heading = document.styles[f"Heading {_SECTION_HEADING_LEVEL}"]
    section_heading.font.name = FONT_NAME
    section_heading.font.color.rgb = ACCENT_COLOR
    section_heading.font.size = SECTION_HEADING_SIZE
    section_heading.font.bold = True
    section_heading.paragraph_format.space_before = Pt(10)
    section_heading.paragraph_format.space_after = Pt(4)

    bullets = document.styles[_BULLET_STYLE]
    bullets.font.name = FONT_NAME
    bullets.font.size = BODY_SIZE
    bullets.font.color.rgb = INK_COLOR
    bullets.paragraph_format.space_after = Pt(2)

    return document


def _add_section_heading(document: DocumentObject, title: str) -> Paragraph:
    """An uppercase, accent-coloured heading over an accent hairline.

    `w:caps` uppercases the rendering only, so the paragraph's text stays
    the ordinary section name for anything parsing the file.
    """
    paragraph = document.add_heading(title, level=_SECTION_HEADING_LEVEL)
    for run in paragraph.runs:
        _style_run(
            run,
            size=SECTION_HEADING_SIZE,
            color=ACCENT_COLOR,
            bold=True,
            caps=True,
        )
    _add_bottom_border(paragraph)
    return paragraph


def _add_tab_stop(document: DocumentObject, paragraph: Paragraph) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        _tab_position(document), WD_TAB_ALIGNMENT.RIGHT
    )


def _add_right_aligned(
    document: DocumentObject, paragraph: Paragraph, text: str
) -> None:
    """Push `text` to the right margin of `paragraph` via a right tab stop."""
    if not text:
        return
    _add_tab_stop(document, paragraph)
    # The tab itself is a run like any other and is formatted like any
    # other — an unstyled one inherits the theme font and shows up as a
    # different face the moment anything measures it.
    _add_text(paragraph, "\t", size=BODY_SIZE, color=MUTED_COLOR)
    _add_text(paragraph, text, size=BODY_SIZE, color=MUTED_COLOR)


def _add_bullets(
    document: DocumentObject,
    bullets: Iterable[str],
    pattern: re.Pattern[str] | None = None,
) -> None:
    """Real list-styled paragraphs, not hand-typed "- " prefixes.

    `pattern`, when given, bolds the profile's own skill terms where they
    already appear in the bullet — see `emphasis_segments` for why that
    cannot alter the text.
    """
    for bullet in _clean_all(bullets):
        paragraph = document.add_paragraph(style=_BULLET_STYLE)
        for fragment, emphasized in emphasis_segments(bullet, pattern):
            _add_text(
                paragraph, fragment, size=BODY_SIZE, color=INK_COLOR, bold=emphasized
            )


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------


def _add_links_line(document: DocumentObject, contact: ContactInfo) -> None:
    """The centered LinkedIn | GitHub | Portfolio line.

    Each label is rendered only when the profile carries that URL, so a
    candidate with no GitHub gets "LinkedIn   |   Portfolio" — never an
    unlinked word, never a placeholder, never a URL guessed from a name.
    The whole paragraph is skipped when none of the three exist.
    """
    links = [
        (LINKEDIN_LABEL, _clean(contact.linkedin_url)),
        (GITHUB_LABEL, _clean(contact.github_url)),
        (PORTFOLIO_LABEL, _clean(contact.portfolio_url)),
    ]
    present = [(label, url) for label, url in links if url]
    if not present:
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    for index, (label, url) in enumerate(present):
        if index:
            _add_text(
                paragraph,
                CONTACT_SEPARATOR,
                size=CONTACT_SIZE,
                color=MUTED_COLOR,
            )
        _style_run(
            add_external_hyperlink(paragraph, label, url),
            size=CONTACT_SIZE,
            color=ACCENT_COLOR,
            underline=True,
        )


def _add_header(document: DocumentObject, profile: ResumeProfile) -> None:
    """Name, optional headline, a contact line, and a links line.

    Written as ordinary body content (heading + paragraphs) rather than a
    Word header/footer, which ATS parsers frequently drop entirely.
    """
    contact = profile.contact
    full_name = _clean(contact.full_name) if contact else ""
    name_paragraph = document.add_heading(
        full_name or NAME_FALLBACK, level=_NAME_HEADING_LEVEL
    )
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in name_paragraph.runs:
        _style_run(run, size=NAME_SIZE, color=INK_COLOR, bold=True)

    headline = _clean(profile.title)
    if headline:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text(
            paragraph, headline, size=CONTACT_SIZE, color=MUTED_COLOR, italic=True
        )

    if contact is None:
        return

    # Location first, matching the reference resume; absent fields drop out
    # entirely rather than leaving a dangling separator behind.
    details = _join(
        [contact.location, contact.phone, contact.email], CONTACT_SEPARATOR
    )
    if details:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text(paragraph, details, size=CONTACT_SIZE, color=MUTED_COLOR)

    _add_links_line(document, contact)


def _add_summary(document: DocumentObject, profile: ResumeProfile) -> None:
    summary = _clean(profile.summary)
    if not summary:
        return
    _add_section_heading(document, SECTION_SUMMARY)
    paragraph = document.add_paragraph()
    _add_text(paragraph, summary, size=BODY_SIZE, color=INK_COLOR)


def _add_skills(document: DocumentObject, profile: ResumeProfile) -> None:
    """One comma-separated line behind a bold lead-in label — how skills read
    on a real resume, and the form keyword-matching ATS parsers handle most
    reliably.
    """
    skills = _clean_all(profile.skills)
    if not skills:
        return
    _add_section_heading(document, SECTION_SKILLS)
    paragraph = document.add_paragraph()
    _add_text(paragraph, SKILLS_LABEL, size=BODY_SIZE, color=INK_COLOR, bold=True)
    _add_text(paragraph, ", ".join(skills), size=BODY_SIZE, color=INK_COLOR)


def _add_work_history(document: DocumentObject, profile: ResumeProfile) -> None:
    """Entries render in the order given — the tailoring step decides the
    emphasis order, so this must not re-sort them.
    """
    entries = profile.work_history
    if not entries:
        return
    _add_section_heading(document, SECTION_EXPERIENCE)
    pattern = _skill_pattern(profile.skills)
    for entry in entries:
        role = document.add_paragraph()
        role.paragraph_format.space_before = Pt(5)
        role.paragraph_format.space_after = Pt(1)
        heading = _join(
            [entry.title, entry.company], TITLE_COMPANY_SEPARATOR
        )
        if heading:
            _add_text(role, heading, size=BODY_SIZE, color=INK_COLOR, bold=True)
        _add_right_aligned(document, role, _date_range(entry))

        _add_bullets(document, entry.bullets, pattern)


def _add_projects(document: DocumentObject, profile: ResumeProfile) -> None:
    projects = _clean_all(profile.projects)
    if not projects:
        return
    _add_section_heading(document, SECTION_PROJECTS)
    _add_bullets(document, projects, _skill_pattern(profile.skills))


def _add_education(document: DocumentObject, profile: ResumeProfile) -> None:
    entries = profile.education
    if not entries:
        return
    _add_section_heading(document, SECTION_EDUCATION)
    for entry in entries:
        qualification = _join(
            [_education_line(entry), _institution_line(entry)],
            TITLE_COMPANY_SEPARATOR,
        )
        if not qualification and not _graduation_year(entry):
            continue
        paragraph = document.add_paragraph()
        if qualification:
            _add_text(
                paragraph, qualification, size=BODY_SIZE, color=INK_COLOR, bold=True
            )
        _add_right_aligned(document, paragraph, _graduation_year(entry))


def _add_certifications(document: DocumentObject, profile: ResumeProfile) -> None:
    certifications = _clean_all(profile.certifications)
    if not certifications:
        return
    _add_section_heading(document, SECTION_CERTIFICATIONS)
    _add_bullets(document, certifications)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def build_resume_document(profile: ResumeProfile) -> bytes:
    """Render `profile` as a complete `.docx` file and return its bytes.

    The returned value is a self-contained Office Open XML document (a zip
    archive) — write it to disk, stream it, or hand it to the tailored-resume
    storage layer as-is. Sections whose source data is empty are omitted
    entirely, so a profile with no certifications produces no
    "Certifications" heading at all.

    Never raises for merely incomplete data: a profile with `contact=None`,
    no work history, no education, or no profile links still produces a
    valid document.
    """
    document = _new_document()

    _add_header(document, profile)
    _add_summary(document, profile)
    _add_skills(document, profile)
    _add_work_history(document, profile)
    _add_projects(document, profile)
    _add_education(document, profile)
    _add_certifications(document, profile)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


__all__ = [
    "ACCENT_COLOR",
    "ACCENT_HEX",
    "CONTACT_SEPARATOR",
    "CURRENT_ROLE_LABEL",
    "FONT_NAME",
    "GITHUB_LABEL",
    "INK_COLOR",
    "LINKEDIN_LABEL",
    "MUTED_COLOR",
    "NAME_FALLBACK",
    "PORTFOLIO_LABEL",
    "SECTION_CERTIFICATIONS",
    "SECTION_EDUCATION",
    "SECTION_EXPERIENCE",
    "SECTION_PROJECTS",
    "SECTION_SKILLS",
    "SECTION_SUMMARY",
    "SKILLS_LABEL",
    "add_external_hyperlink",
    "build_resume_document",
    "emphasis_segments",
]
